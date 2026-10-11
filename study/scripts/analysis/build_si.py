"""Build the Supporting Information package from the Rev 6 figures and tables.

Writes to study/manuscript/supporting_information/:
  SupportingInformation.pdf        contents page, Figures S1 to S8 (the composite PDFs of figures/, placed unchanged), Tables S2 to S5 with their notes
  TableS1_contrast_grid.xlsx       sheet "Table S1" (all contrasts of the planned grid, with estimates where run) and sheet "Key" (codes and counts)
No number is computed here: the figures are placed as they are and the tables are the CSV files of study/manuscript/tables/ (checked by check_table_numbers.py).
Requires PyMuPDF (fitz) and openpyxl.
Usage: python build_si.py
"""
import csv
import html
import tempfile
from pathlib import Path

import fitz
import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

CUR = Path(__file__).resolve().parents[2]
FIG = CUR / "manuscript" / "figures"
TAB = CUR / "manuscript" / "tables"
OUT = CUR / "manuscript" / "supporting_information"
FIGS = [f"S{i}" for i in range(1, 9)]
TABLES = [  # (label, file stem, title)
    ("S2", "TableS2_main_results_P-human", "Main results for the human-only population (P-human)"),
    ("S3", "TableS3_non_primary_pmids", "PMIDs that are not phosphorylation evidence: removed from the publication depth, or kept"),
    ("S4", "TableS4_uniprot_experimental_vs_inferred", "UniProt experimental against inferred modified-residue annotations"),
    ("S5", "TableS5_figure4_sample_sizes", "Sample sizes and estimates behind Figure 4A"),
]
NUMERIC = ("estimate", "ci_low", "ci_high", "second_ci_low", "second_ci_high", "attenuation")
A4 = fitz.paper_rect("a4")
A4L = fitz.paper_rect("a4-l")


def read_csv(path):
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.reader(f))


def build_xlsx(path):
    wb = openpyxl.Workbook()
    wb.properties.creator = ""
    wb.properties.title = "Table S1: contrast grid"
    ws = wb.active
    ws.title = "Table S1"
    rows = read_csv(TAB / "TableS1_contrast_grid.csv")
    head = rows[0]
    ws.append(head)
    for r in rows[1:]:
        ws.append([float(v) if (h in NUMERIC and v != "") else v for h, v in zip(head, r)])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for j, h in enumerate(head, start=1):
        width = max(len(h), *(len(str(ws.cell(row=i, column=j).value or "")) for i in range(2, min(ws.max_row, 80) + 1)))
        ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 8), 60)
    ks = wb.create_sheet("Key")
    for r in read_csv(TAB / "TableS1_key.csv"):
        ks.append(r)
    for c in ks[1]:
        c.font = Font(bold=True)
    ks.freeze_panes = "A2"
    for j, w in enumerate((46, 60, 90, 16), start=1):
        ks.column_dimensions[get_column_letter(j)].width = w
    for row in ks.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return len(rows) - 1, len(read_csv(TAB / "TableS1_key.csv")) - 1


CSS = ("body{font-family:sans-serif;font-size:7.5pt;} h3{font-size:10pt;margin:0 0 6pt 0;} table{border-collapse:collapse;width:100%;} "
       "th{background-color:#e8e8e8;text-align:left;font-weight:bold;border:0.5pt solid #888;padding:2pt;} td{border:0.5pt solid #bbb;padding:2pt;vertical-align:top;} "
       "p{margin:4pt 0 0 0;font-size:7.5pt;}")


def table_html(label, stem, title):
    rows = read_csv(TAB / f"{stem}.csv")
    esc = lambda s: html.escape(s, quote=False)
    out = [f"<html><head><style>{CSS}</style></head><body><h3>Table {label}. {esc(title)}</h3><table><tr>" + "".join(f"<th>{esc(h)}</th>" for h in rows[0]) + "</tr>"]
    for r in rows[1:]:
        out.append("<tr>" + "".join(f"<td>{esc(c)}</td>" for c in r) + "</tr>")
    out.append("</table>")
    notes = TAB / f"{stem}_notes.md"
    if notes.exists():
        for line in notes.read_text(encoding="utf-8").splitlines():
            if line.strip():
                out.append(f"<p>{esc(line.strip())}</p>")
    out.append("</body></html>")
    return "".join(out)


def build_tables(path):
    writer = fitz.DocumentWriter(str(path))
    where = A4L + (36, 36, -36, -50)
    starts = {}
    n = 0
    for label, stem, title in TABLES:
        story = fitz.Story(html=table_html(label, stem, title))
        starts[label] = n
        more = 1
        while more:
            dev = writer.begin_page(A4L)
            more, _ = story.place(where)
            story.draw(dev)
            writer.end_page()
            n += 1
    writer.close()
    return starts, n


def build(out_pdf, out_xlsx):
    OUT.mkdir(parents=True, exist_ok=True)
    n_rows, n_key = build_xlsx(out_xlsx)
    with tempfile.TemporaryDirectory() as tmp:
        tpdf = Path(tmp) / "tables.pdf"
        starts, n_tab_pages = build_tables(tpdf)
        tdoc = fitz.open(tpdf)
        doc = fitz.open()
        doc.new_page(width=A4.width, height=A4.height)  # contents, written last
        fig_page = {}
        for f in FIGS:
            src = fitz.open(FIG / f"Figure{f}.pdf")
            h = src[0].rect.height
            w = src[0].rect.width
            page = doc.new_page(width=A4.width, height=A4.height)
            x0 = (A4.width - w) / 2
            page.insert_text((x0, 48), f"Figure {f}", fontname="hebo", fontsize=11)
            page.show_pdf_page(fitz.Rect(x0, 58, x0 + w, 58 + h), src, 0)
            fig_page[f] = len(doc)
        tab_page = {}
        base = len(doc)
        doc.insert_pdf(tdoc)
        for label, k in starts.items():
            tab_page[label] = base + k + 1
        total = len(doc)
        # contents
        cp = doc[0]
        cp.insert_text((56, 70), "Supporting Information", fontname="hebo", fontsize=14)
        y = 104
        items = [("Table S1", f"Contrast grid ({n_rows} contrasts) and key: separate file TableS1_contrast_grid.xlsx", None)]
        items += [(f"Figure {f}", d, fig_page[f]) for f, d in zip(FIGS, FIG_TITLES)]
        items += [(f"Table {lab}", ttl, tab_page[lab]) for lab, _, ttl in TABLES]
        toc = []
        for label, desc, pg in items:
            cp.insert_text((56, y), label, fontname="hebo", fontsize=9.5)
            rect = fitz.Rect(120, y - 9.5, 500, y + 22)
            cp.insert_textbox(rect, desc, fontname="helv", fontsize=9.5)
            if pg:
                cp.insert_text((520, y), str(pg), fontname="helv", fontsize=9.5)
                toc.append([1, label, pg])
            y += 30 if len(desc) > 70 else 20
        toc.insert(0, [1, "Contents", 1])
        doc.set_toc(toc)
        for i, page in enumerate(doc, start=1):
            r = page.rect
            page.insert_text((36 if r.width > r.height else 56, r.height - 28), "Supporting Information", fontname="helv", fontsize=8)
            page.insert_text((r.width - (36 if r.width > r.height else 56) - 30, r.height - 28), f"{i} / {total}", fontname="helv", fontsize=8)
        doc.set_metadata({"title": "Supporting Information", "author": "", "subject": "", "keywords": ""})
        doc.save(str(out_pdf), garbage=4, deflate=True)
        doc.close()
        tdoc.close()
    return total, fig_page, tab_page, n_rows, n_key


FIG_TITLES = [
    "Database agreement score, study-cluster depth, invalid PMIDs and non-canonical residues",
    "V50 and secondary study-effort index by species and tier; domain annotation against the index",
    "Covariate balance before and after matching on the V50 decile",
    "Comparison of interval methods (site, protein cluster, gene cluster, accession cluster, two-way)",
    "Cluster split and temporal split: partition sizes, test sets and study-cluster counts",
    "Robustness of the GOLD minus BRONZE differences (reviewed accessions, leave-one-cluster-out, tier thresholds)",
    "Database-lookup motif (descriptive only, not graded)",
    "Mean disorder score by tier and source set",
]

if __name__ == "__main__":
    total, fp, tp, nr, nk = build(OUT / "SupportingInformation.pdf", OUT / "TableS1_contrast_grid.xlsx")
    print(f"SupportingInformation.pdf: {total} pages; figures on pages {fp}; tables on pages {tp}; TableS1_contrast_grid.xlsx: {nr} rows, key sheet {nk} rows")
